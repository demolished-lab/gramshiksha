import io
import uuid

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app


@pytest.fixture(name="client")
def client_fixture():
    with TestClient(app) as c:
        yield c


@pytest.fixture(name="headers")
def headers_fixture(client):
    response = client.post("/auth/register", json={
        "email": f"studio.extract.{uuid.uuid4().hex[:8]}@example.com",
        "name": "Studio Importer",
        "password": "Passw0rd@1",
        "role": "student",
    })
    assert response.status_code == 201, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def upload(client, headers, name, data, content_type="application/octet-stream", path="/studio/extract"):
    return client.post(path, headers=headers, files={"file": (name, data, content_type)})


def test_studio_extract_requires_a_session(client):
    response = upload(client, {}, "notes.txt", b"Some notes that are long enough to import.")
    assert response.status_code == 401


def test_studio_extract_returns_clean_text(client, headers):
    body = (
        "Photosynthesis\n\n"
        "  Green plants   use sunlight, water and carbon dioxide.  \n"
        "Chlorophyll absorbs light.\n"
    ).encode()
    response = upload(client, headers, "notes.txt", body, "text/plain")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["filename"] == "notes.txt"
    assert body["text"].startswith("Photosynthesis\nGreen plants use sunlight")
    assert body["truncated"] is False
    # The route-coverage canary reads literal `client.<verb>(...)` calls, so
    # name both mounts explicitly instead of hiding them behind a helper.
    direct = client.post("/studio/extract", headers=headers,
                         files={"file": ("notes.txt", body_text(), "text/plain")})
    assert direct.status_code == 200, direct.text
    # The production-prefixed route is the one the SPA calls.
    prefixed = client.post("/api/studio/extract", headers=headers,
                           files={"file": ("notes.txt", body_text(), "text/plain")})
    assert prefixed.status_code == 200, prefixed.text


def body_text():
    return b"A safe note with enough readable text for the importer to accept it."


def test_studio_extract_strips_scripts_from_html(client, headers):
    html = b"""
    <html><head><style>body{color:red}</style><script>alert(1)</script></head>
    <body><h1>Water cycle</h1><p>Water evaporates and forms clouds.</p>
    <script>steal()</script><p>Rain falls back to the ground.</p></body></html>
    """
    response = upload(client, headers, "notes.html", html, "text/html")
    assert response.status_code == 200, response.text
    text = response.json()["text"]
    assert "Water cycle" in text
    assert "Water evaporates and forms clouds." in text
    assert "alert" not in text and "steal" not in text and "color:red" not in text


def test_studio_extract_reads_docx(client, headers):
    doc = Document()
    doc.add_heading("Photosynthesis", level=1)
    doc.add_paragraph("Plants use sunlight, water and carbon dioxide to make food.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Chlorophyll"
    table.rows[0].cells[1].text = "Absorbs light"
    buffer = io.BytesIO()
    doc.save(buffer)
    response = upload(
        client, headers, "notes.docx", buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert response.status_code == 200, response.text
    text = response.json()["text"]
    assert "Photosynthesis" in text
    assert "Plants use sunlight" in text
    assert "Chlorophyll | Absorbs light" in text


def test_studio_extract_rejects_unsupported_and_empty_files(client, headers):
    unsupported = upload(client, headers, "notes.exe", b"MZ" + b"x" * 40)
    assert unsupported.status_code == 415
    empty = upload(client, headers, "notes.txt", b"")
    assert empty.status_code == 422
    too_short = upload(client, headers, "notes.txt", b"tiny")
    assert too_short.status_code == 422


def test_studio_extract_rejects_a_broken_pdf(client, headers):
    response = upload(client, headers, "notes.pdf", b"%PDF-1.4 not really a pdf")
    assert response.status_code == 422
    assert "PDF" in response.json()["detail"]
