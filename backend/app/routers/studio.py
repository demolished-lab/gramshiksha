"""Text extraction for Study Studio imports.

The learner exports a notebook from NotebookLM as a PDF, DOCX, HTML or text
file, then uploads it to GramShiksha. Parsing happens on the backend because
every phone then gets the same result without downloading a 1 MB PDF library.

Nothing is stored: the upload lives in memory for one request, the extracted
text goes straight back to the browser, and Studio keeps the resulting pack on
the device. Only the first ``MAX_CHARS`` characters are returned; the rest is
reported as truncated instead of being silently dropped.
"""

from __future__ import annotations

import io
import re
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pypdf import PdfReader

from ..ratelimit import rate_limit
from ..security import require_any

router = APIRouter(prefix="/studio", tags=["studio"])

MAX_BYTES = 10 * 1024 * 1024
MAX_CHARS = 120_000
MAX_PDF_PAGES = 80
ALLOWED_EXT = {".txt", ".md", ".markdown", ".html", ".htm", ".pdf", ".docx"}
TEXT_EXT = {".txt", ".md", ".markdown"}


class _HtmlText(HTMLParser):
    """Collect visible HTML text, skipping scripts, styles and SVG data."""

    _BLOCKS = {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6",
               "tr", "section", "article", "blockquote", "pre"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag: str, attrs) -> None:  # noqa: ARG002
        if tag in ("script", "style", "svg", "noscript"):
            self.skip += 1
        elif tag in self._BLOCKS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style", "svg", "noscript"):
            self.skip = max(0, self.skip - 1)
        elif tag in self._BLOCKS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.skip:
            self.parts.append(data)


def _clean_lines(text: str) -> str:
    lines = []
    for line in text.replace("\u00a0", " ").splitlines():
        line = re.sub(r"[ \t]+", " ", line).strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # noqa: BLE001 - pypdf raises several types
                raise ValueError("This PDF is password-protected.") from exc
        pages: list[str] = []
        for page in reader.pages[:MAX_PDF_PAGES]:
            try:
                text = page.extract_text() or ""
            except Exception:  # noqa: BLE001 - one bad page must not kill the import
                continue
            if text.strip():
                pages.append(text)
        return "\n".join(pages)
    except ValueError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ValueError("This PDF could not be read. Try another export.") from exc


def _docx_text(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ValueError("This DOCX file could not be read. Try another export.") from exc
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _html_text(data: bytes) -> str:
    parser = _HtmlText()
    parser.feed(_decode_text(data))
    parser.close()
    return "".join(parser.parts)


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in ALLOWED_EXT:
        raise HTTPException(415, "Use a PDF, DOCX, HTML, Markdown or text file.")
    try:
        if suffix == ".pdf":
            text = _pdf_text(data)
        elif suffix == ".docx":
            text = _docx_text(data)
        elif suffix in (".html", ".htm"):
            text = _html_text(data)
        else:
            text = _decode_text(data)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    return _clean_lines(text)


@router.post("/extract")
async def extract_studio_file(
    file: UploadFile = File(...),
    _rl: None = Depends(rate_limit("studio.extract", 20)),
    _user=Depends(require_any),
):
    data = await file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise HTTPException(413, "File is larger than 10 MB.")
    if not data:
        raise HTTPException(422, "The file is empty.")
    text = extract_text(file.filename or "", data)
    if len(text) < 20:
        raise HTTPException(
            422,
            "No readable text was found. If this is a scanned PDF, export it "
            "as text or copy the notes instead.",
        )
    truncated = len(text) > MAX_CHARS
    return {
        "filename": Path(file.filename or "import").name[:160],
        "text": text[:MAX_CHARS],
        "truncated": truncated,
    }
